"""
Excel (.xlsx) de reportes (punto 28): una hoja con el título, los
filtros con que se generó, el encabezado con los colores de marca, las
filas con su formato (fechas, cantidades, montos) y, si aplica, una
fila de totales. El encabezado queda fijo al desplazarse y con filtros
de Excel, para ordenar y filtrar sin tocar el sistema.

Lo usan Movimientos y "Productos que más se mueven".
"""

from dataclasses import dataclass
from datetime import datetime
from io import BytesIO

from django.http import FileResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

# Mismos colores de marca que los PDF (ver apps/documents/pdf.py).
BRAND_BLUE = "075184"
BRAND_BLUE_SOFT = "EAF3F8"
MUTED_TEXT = "5B6B76"

DATETIME_FORMAT = "dd/mm/yyyy hh:mm"
DATE_FORMAT = "dd/mm/yyyy"
INTEGER_FORMAT = "#,##0"
SIGNED_INTEGER_FORMAT = "+#,##0;-#,##0;0"
MONEY_FORMAT = "#,##0.00"


@dataclass(frozen=True)
class XlsxColumn:
    label: str
    # Ancho en caracteres.
    width: float
    number_format: str | None = None


def build_report_xlsx(
    *,
    title,
    sheet_title,
    columns,
    rows,
    subtitle_lines=(),
    totals=None,
):
    workbook = Workbook()
    sheet = workbook.active
    # Excel no admite más de 31 caracteres en el nombre de una hoja.
    sheet.title = sheet_title[:31]

    generated_at = timezone.localtime(timezone.now())

    row_number = 1

    title_cell = sheet.cell(row=row_number, column=1, value=title)
    title_cell.font = Font(bold=True, size=14, color=BRAND_BLUE)
    row_number += 1

    for line in (
        *subtitle_lines,
        f"Generado el {generated_at.strftime('%d/%m/%Y %H:%M')}",
    ):
        cell = sheet.cell(row=row_number, column=1, value=line)
        cell.font = Font(color=MUTED_TEXT)
        row_number += 1

    # Una fila libre entre el encabezado del documento y la tabla.
    row_number += 1
    header_row = row_number

    header_fill = PatternFill("solid", start_color=BRAND_BLUE)

    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=header_row, column=index, value=column.label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

        sheet.column_dimensions[get_column_letter(index)].width = column.width

    for values in rows:
        row_number += 1
        _write_row(sheet, row_number, columns, values)

    last_data_row = row_number

    if totals is not None:
        row_number += 1
        totals_fill = PatternFill("solid", start_color=BRAND_BLUE_SOFT)

        _write_row(sheet, row_number, columns, totals)

        for index in range(1, len(columns) + 1):
            cell = sheet.cell(row=row_number, column=index)
            cell.font = Font(bold=True)
            cell.fill = totals_fill

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)

    if rows:
        last_column = get_column_letter(len(columns))
        sheet.auto_filter.ref = f"A{header_row}:{last_column}{last_data_row}"

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    return buffer


def xlsx_response(buffer, filename):
    return FileResponse(
        buffer,
        as_attachment=True,
        filename=filename,
        content_type=XLSX_CONTENT_TYPE,
    )


def _write_row(sheet, row_number, columns, values):
    for index, (column, value) in enumerate(zip(columns, values), start=1):
        cell = sheet.cell(row=row_number, column=index, value=_cell_value(value))

        if column.number_format and value not in (None, ""):
            cell.number_format = column.number_format


def _cell_value(value):
    # Excel no guarda zona horaria: la fecha va en hora de Costa Rica.
    if isinstance(value, datetime) and timezone.is_aware(value):
        return timezone.localtime(value).replace(tzinfo=None)

    return value
