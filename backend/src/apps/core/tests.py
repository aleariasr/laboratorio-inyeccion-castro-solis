from decimal import Decimal
from io import StringIO
from types import SimpleNamespace

from django.contrib.auth.models import AnonymousUser, Group, Permission, User
from django.core.management import call_command
from django.test import RequestFactory, TestCase

from apps.core.permissions import (
    AdministrationPermission,
    CustomersPermission,
    DocumentsPermission,
    InjectorsPermission,
    InventoryCountsPermission,
    LocationsPermission,
    MovementsPermission,
    ProductsPermission,
    PurchasesPermission,
    ReportsPermission,
    SalesPermission,
    ServicesPermission,
    SuppliersPermission,
    TransfersPermission,
    ROLE_ADMIN,
    ROLE_CUSTOMERS,
    ROLE_INVENTORY,
    ROLE_READ_ONLY,
    ROLE_SALES,
)
from apps.core.management.commands.setup_roles import ROLE_PERMISSIONS
from apps.inventory.models import PurchaseItem, PurchaseStatus, SupplierProduct
from apps.inventory.selectors import current_stock
from apps.inventory.services import initial_inventory
from apps.core.money import (
    DiscountType,
    DocumentAmountsError,
    compute_document_amounts,
)
from apps.sales.models import Sale, SaleItem, SaleStatus
from apps.sales.services import freeze_sale_amounts


# Tabla de todas las subclases de ModulePermission que declaran
# add_/change_ (y opcionalmente cancel_), usada por
# ModulePermissionMatrixAllModulesTest para no tener que repetir la
# misma mecánica 12 veces a mano. reports/movements quedan fuera de
# esta tabla porque solo declaran view_<module> (ver
# READ_ONLY_MODULE_SPECS) y no tienen add_/change_/cancel_. documents
# tampoco entra aquí: desde §2.2 (proforma) declara add_documents,
# pero no change_/cancel_ (no hay noción de "editar" o "cancelar" un
# documento generado), así que no encaja en ninguna de las dos formas
# — se prueba aparte en test_documents_has_add_but_no_change_or_cancel.
MODULE_PERMISSION_SPECS = [
    ("products", ProductsPermission, False),
    ("locations", LocationsPermission, False),
    ("suppliers", SuppliersPermission, False),
    ("purchases", PurchasesPermission, True),
    ("inventory_counts", InventoryCountsPermission, True),
    ("sales", SalesPermission, True),
    ("customers", CustomersPermission, False),
    ("injectors", InjectorsPermission, False),
    ("services", ServicesPermission, True),
    ("transfers", TransfersPermission, True),
]

# Módulos que en ModulePermissions.Meta.permissions solo declaran
# view_<module>: no existe add_/change_/cancel_ para ellos, así que
# por diseño nadie (fuera de ADMIN/superuser) puede escribir vía la
# API en estos módulos.
READ_ONLY_MODULE_SPECS = [
    ("reports", ReportsPermission),
    ("movements", MovementsPermission),
]


def _module_permission(codename):
    return Permission.objects.get(
        content_type__app_label="core",
        content_type__model="modulepermissions",
        codename=codename,
    )


class ModulePermissionTest(TestCase):
    """
    Prueba la mecánica genérica de ModulePermission: resolución de
    acción -> codename, bypass de superuser/ADMIN, y que is_staff por
    sí solo ya no da acceso a los módulos de negocio.
    """

    def setUp(self):
        self.factory = RequestFactory()

    def _request(self, user, method="get"):
        request_method = getattr(self.factory, method)
        request = request_method("/test/")
        request.user = user
        return request

    def _view(self, action=None):
        return SimpleNamespace(action=action)

    def test_anonymous_user_is_denied(self):
        request = self._request(AnonymousUser())

        self.assertFalse(
            PurchasesPermission().has_permission(
                request,
                self._view(action="list"),
            )
        )

    def test_superuser_is_allowed_without_any_permission(self):
        user = User.objects.create_superuser(
            username="admin",
            password="12345678",
        )
        request = self._request(user, method="post")

        self.assertTrue(
            PurchasesPermission().has_permission(
                request,
                self._view(action="create"),
            )
        )

    def test_admin_group_is_allowed_without_explicit_permissions(self):
        user = User.objects.create_user(
            username="admin-role",
            password="12345678",
        )
        admin_group, _ = Group.objects.get_or_create(name=ROLE_ADMIN)
        user.groups.add(admin_group)

        request = self._request(user, method="post")

        self.assertTrue(
            ServicesPermission().has_permission(
                request,
                self._view(action="cancel"),
            )
        )

    def test_staff_alone_no_longer_bypasses_module_permissions(self):
        """
        A diferencia del esquema anterior basado en roles, is_staff
        por sí solo ya no da acceso a los módulos de negocio: sirve
        para el panel /admin/ de Django, no para el sistema de
        permisos de la aplicación. Solo superuser o el grupo ADMIN
        otorgan bypass.
        """
        user = User.objects.create_user(
            username="staff-only",
            password="12345678",
            is_staff=True,
        )
        request = self._request(user, method="get")

        self.assertFalse(
            ProductsPermission().has_permission(
                request,
                self._view(action="list"),
            )
        )

    def test_user_with_direct_permission_is_allowed(self):
        user = User.objects.create_user(
            username="custom-viewer",
            password="12345678",
        )
        user.user_permissions.add(_module_permission("view_products"))

        request = self._request(user, method="get")

        self.assertTrue(
            ProductsPermission().has_permission(
                request,
                self._view(action="list"),
            )
        )

    def test_user_without_permission_is_denied(self):
        user = User.objects.create_user(
            username="no-perms",
            password="12345678",
        )
        request = self._request(user, method="get")

        self.assertFalse(
            ProductsPermission().has_permission(
                request,
                self._view(action="list"),
            )
        )

    def test_change_permission_does_not_grant_cancel(self):
        user = User.objects.create_user(
            username="purchases-editor",
            password="12345678",
        )
        user.user_permissions.add(_module_permission("change_purchases"))

        request = self._request(user, method="post")
        permission = PurchasesPermission()

        self.assertTrue(
            permission.has_permission(
                request,
                self._view(action="confirm"),
            )
        )
        self.assertFalse(
            permission.has_permission(
                request,
                self._view(action="cancel"),
            )
        )

    def test_cancel_permission_alone_does_not_grant_change(self):
        user = User.objects.create_user(
            username="purchases-canceller",
            password="12345678",
        )
        user.user_permissions.add(_module_permission("cancel_purchases"))

        request = self._request(user, method="post")
        permission = PurchasesPermission()

        self.assertTrue(
            permission.has_permission(
                request,
                self._view(action="cancel"),
            )
        )
        self.assertFalse(
            permission.has_permission(
                request,
                self._view(action="confirm"),
            )
        )

    def test_read_action_only_requires_view_permission(self):
        """
        "labels" es un POST pero de solo lectura (genera un PDF): con
        el permiso de ver productos alcanza, no hace falta el de
        crear.
        """
        user = User.objects.create_user(
            username="labels-viewer",
            password="12345678",
        )
        user.user_permissions.add(_module_permission("view_products"))

        request = self._request(user, method="post")
        permission = ProductsPermission()

        self.assertTrue(
            permission.has_permission(
                request,
                self._view(action="labels"),
            )
        )
        self.assertFalse(
            permission.has_permission(
                request,
                self._view(action="create"),
            )
        )

    def test_role_group_grants_its_module_permissions(self):
        call_command("setup_roles")

        user = User.objects.create_user(
            username="sales-role",
            password="12345678",
        )
        user.groups.add(Group.objects.get(name=ROLE_SALES))

        request = self._request(user, method="post")

        self.assertTrue(
            SalesPermission().has_permission(
                request,
                self._view(action="create"),
            )
        )
        self.assertFalse(
            PurchasesPermission().has_permission(
                request,
                self._view(action="create"),
            )
        )

    def test_administration_permission_denies_anonymous_user(self):
        request = self._request(AnonymousUser())

        self.assertFalse(
            AdministrationPermission().has_permission(
                request,
                view=None,
            )
        )

    def test_administration_permission_allows_admin_group(self):
        user = User.objects.create_user(
            username="administration-admin-role",
            password="12345678",
        )
        admin_group, _ = Group.objects.get_or_create(name=ROLE_ADMIN)
        user.groups.add(admin_group)

        request = self._request(user)

        self.assertTrue(
            AdministrationPermission().has_permission(
                request,
                view=None,
            )
        )

    def test_administration_permission_allows_staff_user(self):
        user = User.objects.create_user(
            username="administration-staff",
            password="12345678",
            is_staff=True,
        )
        request = self._request(user)

        self.assertTrue(
            AdministrationPermission().has_permission(
                request,
                view=None,
            )
        )

    def test_administration_permission_allows_superuser(self):
        user = User.objects.create_superuser(
            username="administration-superuser",
            password="12345678",
        )
        request = self._request(user)

        self.assertTrue(
            AdministrationPermission().has_permission(
                request,
                view=None,
            )
        )

    def test_administration_permission_denies_non_admin_role(self):
        user = User.objects.create_user(
            username="administration-inventory-role",
            password="12345678",
        )
        inventory_group, _ = Group.objects.get_or_create(name=ROLE_INVENTORY)
        user.groups.add(inventory_group)

        request = self._request(user)

        self.assertFalse(
            AdministrationPermission().has_permission(
                request,
                view=None,
            )
        )

    def test_administration_permission_denies_read_only_role(self):
        user = User.objects.create_user(
            username="administration-readonly-role",
            password="12345678",
        )
        read_only_group, _ = Group.objects.get_or_create(name=ROLE_READ_ONLY)
        user.groups.add(read_only_group)

        request = self._request(user)

        self.assertFalse(
            AdministrationPermission().has_permission(
                request,
                view=None,
            )
        )

class SetupRolesCommandTest(TestCase):
    def test_setup_roles_creates_base_groups(self):
        output = StringIO()

        call_command(
            "setup_roles",
            stdout=output,
        )

        expected_roles = {
            ROLE_ADMIN,
            ROLE_INVENTORY,
            ROLE_SALES,
            ROLE_CUSTOMERS,
            ROLE_READ_ONLY,
        }

        existing_roles = set(
            Group.objects.filter(
                name__in=expected_roles,
            ).values_list(
                "name",
                flat=True,
            )
        )

        self.assertEqual(existing_roles, expected_roles)
        self.assertIn("Roles ready.", output.getvalue())

    def test_setup_roles_is_idempotent(self):
        call_command("setup_roles")
        call_command("setup_roles")

        self.assertEqual(
            Group.objects.filter(
                name__in=[
                    ROLE_ADMIN,
                    ROLE_INVENTORY,
                    ROLE_SALES,
                    ROLE_CUSTOMERS,
                    ROLE_READ_ONLY,
                ]
            ).count(),
            5,
        )

    def test_setup_roles_grants_exact_permission_set_per_role(self):
        """
        Compara lo que setup_roles realmente deja en la base de datos
        contra ROLE_PERMISSIONS. Si alguien agrega/quita un permiso de
        módulo en ModulePermissions.Meta.permissions o en
        ROLE_PERMISSIONS sin actualizar el otro, este test lo detecta
        (permiso de más o permiso de menos), en vez de descubrirlo en
        producción cuando un rol no puede hacer algo que debería, o
        puede hacer algo que no debería.
        """
        call_command("setup_roles")

        all_module_codenames = set(
            Permission.objects.filter(
                content_type__app_label="core",
                content_type__model="modulepermissions",
            ).values_list(
                "codename",
                flat=True,
            )
        )

        # 38 = cantidad de tuplas en ModulePermissions.Meta.permissions
        # al momento de escribir este test. Un cambio en este número
        # es la señal de que hay que revisar ROLE_PERMISSIONS también.
        # (37 → 38 con view_costs, del rediseño de roles del punto 13;
        # 38 → 42 con los traslados, punto 6 del plan del 2026-10-22.)
        self.assertEqual(len(all_module_codenames), 42)

        admin_group = Group.objects.get(name=ROLE_ADMIN)
        admin_codenames = set(
            admin_group.permissions.values_list(
                "codename",
                flat=True,
            )
        )

        # ADMIN no está en ROLE_PERMISSIONS (recibe todo por código en
        # setup_roles), así que se compara aparte contra el catálogo
        # completo.
        self.assertEqual(admin_codenames, all_module_codenames)

        for role, expected_codenames in ROLE_PERMISSIONS.items():
            with self.subTest(role=role):
                group = Group.objects.get(name=role)
                actual_codenames = set(
                    group.permissions.values_list(
                        "codename",
                        flat=True,
                    )
                )

                self.assertEqual(
                    actual_codenames,
                    set(expected_codenames),
                )

    def test_setup_roles_permission_assignment_is_idempotent(self):
        call_command("setup_roles")

        first_run_codenames = {
            role: set(
                Group.objects.get(name=role).permissions.values_list(
                    "codename",
                    flat=True,
                )
            )
            for role in ROLE_PERMISSIONS
        }

        call_command("setup_roles")

        second_run_codenames = {
            role: set(
                Group.objects.get(name=role).permissions.values_list(
                    "codename",
                    flat=True,
                )
            )
            for role in ROLE_PERMISSIONS
        }

        self.assertEqual(first_run_codenames, second_run_codenames)


class ModulePermissionMatrixAllModulesTest(TestCase):
    """
    Repite, para las 12 subclases concretas de ModulePermission, la
    misma mecánica que ModulePermissionTest solo prueba sobre una
    muestra de 3-4 de ellas (view/add/change/cancel). El objetivo es
    que un cambio futuro en un solo módulo (ej. agregar cancel_actions
    a Suppliers, o quitarle read_actions a Products) no pueda romper
    otro módulo sin que quede evidenciado acá.
    """

    def setUp(self):
        self.factory = RequestFactory()

    def _request(self, user, method="get"):
        request_method = getattr(self.factory, method)
        request = request_method("/test/")
        request.user = user
        return request

    def _view(self, action=None):
        return SimpleNamespace(action=action)

    def _user_with_permission(self, codename, username):
        user = User.objects.create_user(
            username=username,
            password="12345678",
        )
        user.user_permissions.add(_module_permission(codename))
        return user

    def test_view_permission_allows_read_and_denies_write_for_every_module(self):
        all_specs = [
            (module, permission_class)
            for module, permission_class, _has_cancel in MODULE_PERMISSION_SPECS
        ] + list(READ_ONLY_MODULE_SPECS)

        for module, permission_class in all_specs:
            with self.subTest(module=module):
                user = self._user_with_permission(
                    f"view_{module}",
                    f"view-only-{module}",
                )
                permission = permission_class()

                read_request = self._request(user, method="get")
                self.assertTrue(
                    permission.has_permission(
                        read_request,
                        self._view(action="list"),
                    )
                )

                create_request = self._request(user, method="post")
                self.assertFalse(
                    permission.has_permission(
                        create_request,
                        self._view(action="create"),
                    )
                )

                update_request = self._request(user, method="patch")
                self.assertFalse(
                    permission.has_permission(
                        update_request,
                        self._view(action="partial_update"),
                    )
                )

    def test_add_permission_allows_create_and_denies_read_and_update(self):
        for module, permission_class, _has_cancel in MODULE_PERMISSION_SPECS:
            with self.subTest(module=module):
                user = self._user_with_permission(
                    f"add_{module}",
                    f"add-only-{module}",
                )
                permission = permission_class()

                create_request = self._request(user, method="post")
                self.assertTrue(
                    permission.has_permission(
                        create_request,
                        self._view(action="create"),
                    )
                )

                read_request = self._request(user, method="get")
                self.assertFalse(
                    permission.has_permission(
                        read_request,
                        self._view(action="list"),
                    )
                )

                update_request = self._request(user, method="patch")
                self.assertFalse(
                    permission.has_permission(
                        update_request,
                        self._view(action="partial_update"),
                    )
                )

    def test_change_permission_allows_update_and_denies_create(self):
        for module, permission_class, _has_cancel in MODULE_PERMISSION_SPECS:
            with self.subTest(module=module):
                user = self._user_with_permission(
                    f"change_{module}",
                    f"change-only-{module}",
                )
                permission = permission_class()

                update_request = self._request(user, method="patch")
                self.assertTrue(
                    permission.has_permission(
                        update_request,
                        self._view(action="partial_update"),
                    )
                )

                create_request = self._request(user, method="post")
                self.assertFalse(
                    permission.has_permission(
                        create_request,
                        self._view(action="create"),
                    )
                )

    def test_change_permission_alone_does_not_grant_cancel(self):
        for module, permission_class, has_cancel in MODULE_PERMISSION_SPECS:
            if not has_cancel:
                continue

            with self.subTest(module=module):
                user = self._user_with_permission(
                    f"change_{module}",
                    f"change-only-cancel-check-{module}",
                )
                permission = permission_class()

                cancel_request = self._request(user, method="post")
                self.assertFalse(
                    permission.has_permission(
                        cancel_request,
                        self._view(action="cancel"),
                    )
                )

    def test_cancel_permission_alone_allows_cancel_and_denies_update(self):
        for module, permission_class, has_cancel in MODULE_PERMISSION_SPECS:
            if not has_cancel:
                continue

            with self.subTest(module=module):
                user = self._user_with_permission(
                    f"cancel_{module}",
                    f"cancel-only-{module}",
                )
                permission = permission_class()

                cancel_request = self._request(user, method="post")
                self.assertTrue(
                    permission.has_permission(
                        cancel_request,
                        self._view(action="cancel"),
                    )
                )

                update_request = self._request(user, method="patch")
                self.assertFalse(
                    permission.has_permission(
                        update_request,
                        self._view(action="partial_update"),
                    )
                )

    def test_read_only_modules_have_no_write_permission_in_database(self):
        """
        reports y movements solo declaran view_<module> en
        ModulePermissions: no existe add_/change_/cancel_ para ellos.
        Si alguna vez se agrega una de esas variantes sin actualizar
        este test, hay que revisar también setup_roles y el frontend,
        porque implicaría que el módulo dejó de ser de solo lectura.
        """
        for module, _permission_class in READ_ONLY_MODULE_SPECS:
            with self.subTest(module=module):
                self.assertFalse(
                    Permission.objects.filter(
                        content_type__app_label="core",
                        content_type__model="modulepermissions",
                        codename=f"add_{module}",
                    ).exists()
                )
                self.assertFalse(
                    Permission.objects.filter(
                        content_type__app_label="core",
                        content_type__model="modulepermissions",
                        codename=f"change_{module}",
                    ).exists()
                )
                self.assertFalse(
                    Permission.objects.filter(
                        content_type__app_label="core",
                        content_type__model="modulepermissions",
                        codename=f"cancel_{module}",
                    ).exists()
                )

    def test_documents_has_add_but_no_change_or_cancel(self):
        """
        documents dejó de ser puramente de solo lectura en §2.2
        (proforma), que agregó add_documents — pero sigue sin
        change_/cancel_documents, porque no existe la noción de
        "editar" o "cancelar" un documento ya generado (proforma o
        factura). DocumentsPermission importado solo para dejar
        explícito qué clase gobierna este módulo.
        """
        self.assertIsNotNone(DocumentsPermission)

        self.assertTrue(
            Permission.objects.filter(
                content_type__app_label="core",
                content_type__model="modulepermissions",
                codename="view_documents",
            ).exists()
        )
        self.assertTrue(
            Permission.objects.filter(
                content_type__app_label="core",
                content_type__model="modulepermissions",
                codename="add_documents",
            ).exists()
        )
        self.assertFalse(
            Permission.objects.filter(
                content_type__app_label="core",
                content_type__model="modulepermissions",
                codename="change_documents",
            ).exists()
        )
        self.assertFalse(
            Permission.objects.filter(
                content_type__app_label="core",
                content_type__model="modulepermissions",
                codename="cancel_documents",
            ).exists()
        )


from datetime import date, timedelta

from django.utils import timezone
from rest_framework import status
from rest_framework.test import APITestCase

from apps.cash.models import CashClosing
from apps.cash.selectors import pending_closing_week_start
from apps.customers.models import (
    Customer,
    Injector,
    InjectorServiceRecord,
    InjectorServiceStatus,
)
from apps.inventory.models import (
    Product,
    Purchase,
    StorageLocation,
    Supplier,
)
from apps.sales.models import Sale, SaleStatus


class UniversalSearchApiTest(APITestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="search-user",
            password="12345678",
        )
        self.client.force_authenticate(self.user)
        call_command("setup_roles")
        self.user.groups.add(Group.objects.get(name=ROLE_ADMIN))

        self.location = StorageLocation.objects.create(
            code="A124",
            description="Estante A posición 124",
            created_by=self.user,
            updated_by=self.user,
        )

        self.product = Product.objects.create(
            standard_code="1-423-124-108",
            name="Tornillo bloqueo Cummins",
            description="Pieza de prueba",
            storage_location=self.location,
            created_by=self.user,
            updated_by=self.user,
        )

        self.supplier = Supplier.objects.create(
            name="Proveedor Central",
            phone="2222-2222",
            email="proveedor@example.com",
            country="Costa Rica",
            created_by=self.user,
            updated_by=self.user,
        )

        self.purchase = Purchase.objects.create(
            supplier=self.supplier,
            invoice_number="FAC-001",
            purchase_date=date.today(),
            currency="CRC",
            created_by=self.user,
            updated_by=self.user,
        )

        self.customer = Customer.objects.create(
            customer_type="PERSON",
            display_name="Cliente Diesel",
            phone="8888-8888",
            email="cliente@example.com",
            identification="1-1111-1111",
            created_by=self.user,
            updated_by=self.user,
        )

        self.injector = Injector.objects.create(
            customer=self.customer,
            injector_number="INY-001",
            description="Inyector Bosch",
            created_by=self.user,
            updated_by=self.user,
        )

        self.sale = Sale.objects.create(
            customer=self.customer,
            sale_date=date.today(),
            status=SaleStatus.CONFIRMED,
            created_by=self.user,
            updated_by=self.user,
        )

        self.service_record = InjectorServiceRecord.objects.create(
            injector=self.injector,
            received_at=timezone.now(),
            created_by=self.user,
            updated_by=self.user,
        )

    def test_search_requires_authentication(self):
        self.client.force_authenticate(user=None)

        response = self.client.get(
            "/api/search/",
            {
                "q": "A124",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_search_short_query_returns_empty_results(self):
        response = self.client.get(
            "/api/search/",
            {
                "q": "A",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["query"], "A")
        self.assertEqual(response.data["results"]["products"], [])
        self.assertEqual(response.data["results"]["locations"], [])

    def test_search_finds_location(self):
        response = self.client.get(
            "/api/search/",
            {
                "q": "A124",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response.data["results"]["locations"][0]["code"],
            "A124",
        )

    def test_search_finds_product_by_standard_code(self):
        response = self.client.get(
            "/api/search/",
            {
                "q": "1-423",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response.data["results"]["products"][0]["standard_code"],
            "1-423-124-108",
        )

    def test_search_by_standard_code_finds_all_variants(self):
        Product.objects.create(
            standard_code="1-423-124-108",
            name="Tornillo bloqueo Cummins (genérico)",
            variant_kind="GENERIC",
            storage_location=self.location,
            created_by=self.user,
            updated_by=self.user,
        )

        response = self.client.get(
            "/api/search/",
            {
                "q": "1-423",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            len(response.data["results"]["products"]),
            2,
        )

    def test_search_finds_supplier(self):
        response = self.client.get(
            "/api/search/",
            {
                "q": "Central",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response.data["results"]["suppliers"][0]["name"],
            "PROVEEDOR CENTRAL",
        )

    def test_search_finds_purchase(self):
        response = self.client.get(
            "/api/search/",
            {
                "q": "FAC-001",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response.data["results"]["purchases"][0]["invoice_number"],
            "FAC-001",
        )

    def test_search_finds_customer(self):
        response = self.client.get(
            "/api/search/",
            {
                "q": "Diesel",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response.data["results"]["customers"][0]["display_name"],
            "CLIENTE DIESEL",
        )

    def test_search_finds_injector(self):
        response = self.client.get(
            "/api/search/",
            {
                "q": "INY-001",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response.data["results"]["injectors"][0]["injector_number"],
            "INY-001",
        )

    def test_search_finds_product_by_name(self):
        response = self.client.get(
            "/api/search/",
            {
                "q": "Tornillo",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response.data["results"]["products"][0]["standard_code"],
            "1-423-124-108",
        )

    def test_search_finds_product_by_description(self):
        response = self.client.get(
            "/api/search/",
            {
                "q": "Pieza de prueba",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response.data["results"]["products"][0]["standard_code"],
            "1-423-124-108",
        )

    def test_search_finds_customer_by_identification(self):
        response = self.client.get(
            "/api/search/",
            {
                "q": "1-1111-1111",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response.data["results"]["customers"][0]["display_name"],
            "CLIENTE DIESEL",
        )

    def test_search_finds_customer_by_phone(self):
        response = self.client.get(
            "/api/search/",
            {
                "q": "8888-8888",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response.data["results"]["customers"][0]["display_name"],
            "CLIENTE DIESEL",
        )

    def test_search_finds_sale_by_customer_name(self):
        response = self.client.get(
            "/api/search/",
            {
                "q": "Diesel",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response.data["results"]["sales"][0]["id"],
            self.sale.id,
        )

    def test_search_finds_service_record_by_injector_number(self):
        response = self.client.get(
            "/api/search/",
            {
                "q": "INY-001",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response.data["results"]["services"][0]["id"],
            self.service_record.id,
        )

    def test_search_finds_service_record_by_customer_name(self):
        response = self.client.get(
            "/api/search/",
            {
                "q": "Diesel",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response.data["results"]["services"][0]["id"],
            self.service_record.id,
        )

    def test_search_hides_categories_the_user_cannot_read(self):
        limited_user = User.objects.create_user(
            username="search-limited",
            password="12345678",
        )
        limited_user.user_permissions.add(
            Permission.objects.get(
                content_type__app_label="core",
                content_type__model="modulepermissions",
                codename="view_products",
            ),
        )
        self.client.force_authenticate(limited_user)

        response = self.client.get(
            "/api/search/",
            {"q": "Diesel"},
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["results"]["customers"], [])

    def test_search_shows_category_the_user_can_read(self):
        limited_user = User.objects.create_user(
            username="search-customers-only",
            password="12345678",
        )
        limited_user.user_permissions.add(
            Permission.objects.get(
                content_type__app_label="core",
                content_type__model="modulepermissions",
                codename="view_customers",
            ),
        )
        self.client.force_authenticate(limited_user)

        response = self.client.get(
            "/api/search/",
            {"q": "Diesel"},
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response.data["results"]["customers"][0]["display_name"],
            "CLIENTE DIESEL",
        )
        self.assertEqual(response.data["results"]["products"], [])

    def test_search_hides_sales_and_services_without_permission(self):
        limited_user = User.objects.create_user(
            username="search-no-sales-services",
            password="12345678",
        )
        limited_user.user_permissions.add(
            Permission.objects.get(
                content_type__app_label="core",
                content_type__model="modulepermissions",
                codename="view_customers",
            ),
        )
        self.client.force_authenticate(limited_user)

        response = self.client.get(
            "/api/search/",
            {"q": "Diesel"},
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["results"]["sales"], [])
        self.assertEqual(response.data["results"]["services"], [])


class DashboardSummaryApiTest(APITestCase):
    def setUp(self):
        call_command("setup_roles")

        self.user = User.objects.create_user(
            username="dashboard-admin",
            password="12345678",
        )
        self.user.groups.add(Group.objects.get(name=ROLE_ADMIN))
        self.client.force_authenticate(self.user)

        self.location = StorageLocation.objects.create(
            code="D101",
            created_by=self.user,
            updated_by=self.user,
        )

        self.low_stock_product = Product.objects.create(
            standard_code="LOW-001",
            name="Producto bajo mínimo",
            storage_location=self.location,
            minimum_stock=5,
            created_by=self.user,
            updated_by=self.user,
        )

        self.customer = Customer.objects.create(
            customer_type="PERSON",
            display_name="Cliente Dashboard",
            identification="2-2222-2222",
            created_by=self.user,
            updated_by=self.user,
        )

        self.injector = Injector.objects.create(
            customer=self.customer,
            injector_number="INY-900",
            created_by=self.user,
            updated_by=self.user,
        )

        InjectorServiceRecord.objects.create(
            injector=self.injector,
            received_at=timezone.now(),
            status=InjectorServiceStatus.IN_PROGRESS,
            created_by=self.user,
            updated_by=self.user,
        )
        self.ready_service = InjectorServiceRecord.objects.create(
            injector=self.injector,
            received_at=timezone.now(),
            status=InjectorServiceStatus.READY,
            created_by=self.user,
            updated_by=self.user,
        )

        self.draft_sale = Sale.objects.create(
            customer=self.customer,
            sale_date=date.today(),
            status=SaleStatus.DRAFT,
            created_by=self.user,
            updated_by=self.user,
        )

        self.supplier = Supplier.objects.create(
            name="Proveedor Dashboard",
            created_by=self.user,
            updated_by=self.user,
        )

        self.draft_purchase = Purchase.objects.create(
            supplier=self.supplier,
            invoice_number="FAC-900",
            purchase_date=date.today(),
            currency="CRC",
            status=PurchaseStatus.DRAFT,
            created_by=self.user,
            updated_by=self.user,
        )

    def test_dashboard_summary_returns_all_sections_for_admin(self):
        response = self.client.get("/api/dashboard/summary/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)

        self.assertEqual(
            response.data["low_stock_products"][0]["standard_code"],
            "LOW-001",
        )
        self.assertEqual(response.data["services_in_progress_count"], 1)
        self.assertEqual(
            response.data["services_ready"][0]["id"],
            self.ready_service.id,
        )
        self.assertEqual(
            response.data["services_ready"][0]["injector_number"],
            "INY-900",
        )
        self.assertEqual(len(response.data["draft_sales"]), 1)
        self.assertEqual(
            response.data["draft_sales"][0]["id"],
            self.draft_sale.id,
        )
        self.assertEqual(len(response.data["draft_purchases"]), 1)
        self.assertEqual(
            response.data["draft_purchases"][0]["invoice_number"],
            "FAC-900",
        )
        self.assertIsNotNone(response.data["cash_pending_week_start"])

    def test_dashboard_summary_cash_pending_is_null_once_closed(self):
        pending_week_start = pending_closing_week_start()

        CashClosing.objects.create(
            week_start=pending_week_start,
            week_end=pending_week_start + timedelta(days=6),
            expected_total=Decimal("0"),
            expected_cash=Decimal("0"),
            expected_card=Decimal("0"),
            expected_transfer=Decimal("0"),
            expected_other=Decimal("0"),
            counted_total=Decimal("0"),
            difference=Decimal("0"),
            created_by=self.user,
            updated_by=self.user,
        )

        response = self.client.get("/api/dashboard/summary/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIsNone(response.data["cash_pending_week_start"])

    def test_dashboard_summary_hides_sections_without_permission(self):
        limited_user = User.objects.create_user(
            username="dashboard-sales-only",
            password="12345678",
        )
        limited_user.groups.add(Group.objects.get(name=ROLE_SALES))
        self.client.force_authenticate(limited_user)

        response = self.client.get("/api/dashboard/summary/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIsNotNone(response.data["draft_sales"])
        self.assertIsNone(response.data["low_stock_products"])
        self.assertIsNone(response.data["services_in_progress_count"])
        self.assertIsNone(response.data["services_ready"])
        self.assertIsNone(response.data["draft_purchases"])
        self.assertIsNone(response.data["cash_pending_week_start"])

    def test_low_stock_card_follows_full_inventory_module(self):
        # Ventas y Clientes tienen view_products para buscar productos
        # en sus formularios, pero la alerta de bajo mínimo es para los
        # roles con el módulo de inventario completo.
        expectations = {
            ROLE_INVENTORY: True,
            ROLE_READ_ONLY: True,
            ROLE_SALES: False,
            ROLE_CUSTOMERS: False,
        }

        for role, sees_card in expectations.items():
            with self.subTest(role=role):
                user = User.objects.create_user(
                    username=f"dashboard-low-stock-{role.lower()}",
                    password="12345678",
                )
                user.groups.add(Group.objects.get(name=role))
                self.client.force_authenticate(user)

                response = self.client.get("/api/dashboard/summary/")

                self.assertEqual(response.status_code, status.HTTP_200_OK)
                self.assertEqual(
                    response.data["low_stock_products"] is not None,
                    sees_card,
                )

    def test_dashboard_summary_requires_authentication(self):
        self.client.force_authenticate(user=None)

        response = self.client.get("/api/dashboard/summary/")

        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)


class InventoryReportsApiTest(APITestCase):
    def setUp(self):
        call_command("setup_roles")

        self.inventory_user = User.objects.create_user(
            username="reports-inventory",
            password="12345678",
        )
        self.inventory_user.groups.add(
            Group.objects.get(name=ROLE_INVENTORY),
        )

        self.read_only_user = User.objects.create_user(
            username="reports-readonly",
            password="12345678",
        )
        self.read_only_user.groups.add(
            Group.objects.get(name=ROLE_READ_ONLY),
        )

        self.plain_user = User.objects.create_user(
            username="reports-plain",
            password="12345678",
        )

        self.admin_user = User.objects.create_user(
            username="reports-admin",
            password="12345678",
        )
        self.admin_user.groups.add(
            Group.objects.get(name=ROLE_ADMIN),
        )

        self.location = StorageLocation.objects.create(
            code="B200",
            description="Estante B",
            created_by=self.inventory_user,
            updated_by=self.inventory_user,
        )

        self.product = Product.objects.create(
            standard_code="REP-001",
            name="Producto reporte",
            description="Producto para reportes",
            storage_location=self.location,
            minimum_stock=5,
            created_by=self.inventory_user,
            updated_by=self.inventory_user,
        )

        self.product_with_stock = Product.objects.create(
            standard_code="REP-002",
            name="Producto con stock",
            description="Producto con inventario",
            storage_location=self.location,
            minimum_stock=1,
            created_by=self.inventory_user,
            updated_by=self.inventory_user,
        )

        initial_inventory(
            product=self.product_with_stock,
            quantity=10,
            user=self.inventory_user,
        )

    def test_low_stock_report_lists_products_below_minimum(self):
        self.client.force_authenticate(self.admin_user)

        response = self.client.get(
            "/api/reports/low-stock-products/"
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)

        standard_codes = {
            item["standard_code"]
            for item in response.data["results"]
        }

        self.assertIn("REP-001", standard_codes)
        self.assertNotIn("REP-002", standard_codes)

    def test_stock_by_location_report_groups_products(self):
        self.client.force_authenticate(self.admin_user)

        response = self.client.get(
            "/api/reports/stock-by-location/"
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["results"][0]["code"], "B200")

        product_codes = {
            item["standard_code"]
            for item in response.data["results"][0]["products"]
        }

        self.assertIn("REP-001", product_codes)
        self.assertIn("REP-002", product_codes)

    def test_product_movements_report_lists_movements(self):
        self.client.force_authenticate(self.admin_user)

        response = self.client.get(
            "/api/reports/product-movements/",
            {
                "product": self.product_with_stock.id,
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response.data["product"]["standard_code"],
            "REP-002",
        )
        self.assertEqual(len(response.data["results"]), 1)
        self.assertEqual(current_stock(self.product_with_stock), 10)

    def test_product_movements_report_requires_product(self):
        self.client.force_authenticate(self.admin_user)

        response = self.client.get(
            "/api/reports/product-movements/"
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_reports_block_read_only_user(self):
        # Los reportes ahora son exclusivos de ADMIN: READ_ONLY ya no
        # incluye view_reports en ROLE_PERMISSIONS (setup_roles.py).
        self.client.force_authenticate(self.read_only_user)

        response = self.client.get(
            "/api/reports/low-stock-products/"
        )

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_reports_block_inventory_only_user(self):
        # INVENTORY perdió view_reports por el mismo motivo.
        self.client.force_authenticate(self.inventory_user)

        response = self.client.get(
            "/api/reports/low-stock-products/"
        )

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_reports_allow_admin_user_to_read(self):
        self.client.force_authenticate(self.admin_user)

        response = self.client.get(
            "/api/reports/low-stock-products/"
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)

    def test_reports_block_authenticated_user_without_group(self):
        self.client.force_authenticate(self.plain_user)

        response = self.client.get(
            "/api/reports/low-stock-products/"
        )

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_reports_require_authentication(self):
        response = self.client.get(
            "/api/reports/low-stock-products/"
        )

        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_reports_block_sales_only_user(self):
        # SALES perdió view_reports por el mismo motivo.
        sales_user = User.objects.create_user(
            username="reports-sales-only",
            password="12345678",
        )

        sales_user.groups.add(
            Group.objects.get(name=ROLE_SALES),
        )

        self.client.force_authenticate(sales_user)

        response = self.client.get(
            "/api/reports/low-stock-products/"
        )

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_low_stock_report_is_paginated(self):
        Product.objects.create(
            standard_code="REP-003",
            name="Segundo producto bajo mínimo",
            description="Otro producto para probar paginación",
            storage_location=self.location,
            minimum_stock=5,
            created_by=self.inventory_user,
            updated_by=self.inventory_user,
        )

        self.client.force_authenticate(self.admin_user)

        first_page = self.client.get(
            "/api/reports/low-stock-products/",
            {"page_size": 1},
        )

        self.assertEqual(first_page.status_code, status.HTTP_200_OK)
        self.assertEqual(first_page.data["count"], 2)
        self.assertEqual(len(first_page.data["results"]), 1)
        self.assertIsNotNone(first_page.data["next"])
        self.assertIsNone(first_page.data["previous"])

        second_page = self.client.get(
            "/api/reports/low-stock-products/",
            {"page_size": 1, "page": 2},
        )

        self.assertEqual(second_page.status_code, status.HTTP_200_OK)
        self.assertEqual(len(second_page.data["results"]), 1)
        self.assertIsNone(second_page.data["next"])
        self.assertIsNotNone(second_page.data["previous"])

    def test_stock_by_location_report_is_paginated(self):
        second_location = StorageLocation.objects.create(
            code="B201",
            description="Estante B2",
            created_by=self.inventory_user,
            updated_by=self.inventory_user,
        )

        Product.objects.create(
            standard_code="REP-004",
            name="Producto en otra ubicación",
            storage_location=second_location,
            created_by=self.inventory_user,
            updated_by=self.inventory_user,
        )

        self.client.force_authenticate(self.admin_user)

        first_page = self.client.get(
            "/api/reports/stock-by-location/",
            {"page_size": 1},
        )

        self.assertEqual(first_page.status_code, status.HTTP_200_OK)
        self.assertEqual(first_page.data["count"], 2)
        self.assertEqual(len(first_page.data["results"]), 1)
        self.assertIsNotNone(first_page.data["next"])
        self.assertIsNone(first_page.data["previous"])

        second_page = self.client.get(
            "/api/reports/stock-by-location/",
            {"page_size": 1, "page": 2},
        )

        self.assertEqual(second_page.status_code, status.HTTP_200_OK)
        self.assertEqual(len(second_page.data["results"]), 1)
        self.assertIsNone(second_page.data["next"])
        self.assertIsNotNone(second_page.data["previous"])

    def test_stock_by_location_report_filters_by_exact_location(self):
        second_location = StorageLocation.objects.create(
            code="B202",
            description="Estante B3",
            created_by=self.inventory_user,
            updated_by=self.inventory_user,
        )

        Product.objects.create(
            standard_code="REP-005",
            name="Producto en tercera ubicación",
            storage_location=second_location,
            created_by=self.inventory_user,
            updated_by=self.inventory_user,
        )

        self.client.force_authenticate(self.admin_user)

        response = self.client.get(
            "/api/reports/stock-by-location/",
            {"location": self.location.id},
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data["results"]), 1)
        self.assertEqual(response.data["results"][0]["code"], "B200")

    def test_stock_by_location_report_rejects_invalid_location(self):
        self.client.force_authenticate(self.admin_user)

        response = self.client.get(
            "/api/reports/stock-by-location/",
            {"location": "not-a-number"},
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_stock_by_location_report_filters_by_product_search(self):
        self.client.force_authenticate(self.admin_user)

        response = self.client.get(
            "/api/reports/stock-by-location/",
            {"q": "REP-002"},
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data["results"]), 1)

        product_codes = {
            item["standard_code"]
            for item in response.data["results"][0]["products"]
        }

        self.assertEqual(product_codes, {"REP-002"})

class BusinessReportsApiTest(APITestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="business-reports",
            password="12345678",
        )

        call_command("setup_roles")

        # Los reportes de negocio ahora son exclusivos de ADMIN: ver
        # test_business_reports_block_sales_only_user y
        # test_business_reports_block_inventory_only_user más abajo
        # para la verificación del bloqueo a roles no administrativos.
        self.user.groups.add(
            Group.objects.get(name=ROLE_ADMIN),
        )

        self.client.force_authenticate(self.user)

        self.location = StorageLocation.objects.create(
            code="C300",
            description="Estante C",
            created_by=self.user,
            updated_by=self.user,
        )

        self.product = Product.objects.create(
            standard_code="TOP-001",
            name="Producto vendido",
            storage_location=self.location,
            created_by=self.user,
            updated_by=self.user,
        )

        self.supplier = Supplier.objects.create(
            name="Proveedor Reportes",
            created_by=self.user,
            updated_by=self.user,
        )

        self.supplier_product = SupplierProduct.objects.create(
            supplier=self.supplier,
            product=self.product,
            created_by=self.user,
            updated_by=self.user,
        )

        self.purchase = Purchase.objects.create(
            supplier=self.supplier,
            invoice_number="REP-FAC-001",
            purchase_date=date(2026, 7, 1),
            currency="CRC",
            status=PurchaseStatus.CONFIRMED,
            created_by=self.user,
            updated_by=self.user,
        )

        PurchaseItem.objects.create(
            purchase=self.purchase,
            supplier_product=self.supplier_product,
            quantity=4,
            unit_cost=Decimal("100.0000"),
            created_by=self.user,
            updated_by=self.user,
        )

        self.customer = Customer.objects.create(
            display_name="Cliente Reportes",
            phone="8888-1111",
            created_by=self.user,
            updated_by=self.user,
        )

        self.sale = Sale.objects.create(
            customer=self.customer,
            sale_date=date(2026, 7, 2),
            currency="CRC",
            status=SaleStatus.CONFIRMED,
            created_by=self.user,
            updated_by=self.user,
        )

        SaleItem.objects.create(
            sale=self.sale,
            product=self.product,
            quantity=3,
            unit_price=Decimal("250.0000"),
            created_by=self.user,
            updated_by=self.user,
        )

    def test_purchases_by_supplier_report(self):
        response = self.client.get(
            "/api/reports/purchases-by-supplier/",
            {
                "date_from": "2026-07-01",
                "date_to": "2026-07-31",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data["results"]), 1)

        item = response.data["results"][0]

        self.assertEqual(
            item["supplier"]["name"],
            "PROVEEDOR REPORTES",
        )
        self.assertEqual(item["purchase_count"], 1)
        self.assertEqual(
            Decimal(item["invoice_subtotal"]),
            Decimal("400.0000"),
        )

    def test_sales_summary_report_by_day(self):
        # Reemplaza al viejo "Ventas por fecha" (punto 20): agrupado por
        # día da lo mismo que antes, más el desglose.
        response = self.client.get(
            "/api/reports/sales-summary/",
            {
                "date_from": "2026-07-01",
                "date_to": "2026-07-31",
                "group_by": "day",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data["periods"]), 1)

        period = response.data["periods"][0]

        self.assertEqual(period["period_start"], date(2026, 7, 2))
        self.assertEqual(period["period_end"], date(2026, 7, 2))
        self.assertEqual(period["sales"]["count"], 1)
        self.assertEqual(period["sales"]["total"], Decimal("750.00"))
        self.assertEqual(period["services"]["count"], 0)
        self.assertEqual(period["total"]["total"], Decimal("750.00"))
        self.assertIsNone(response.data["comparison"])

    def test_sales_reports_use_amount_charged_but_top_products_stay_gross(self):
        # Punto 19: 750 − 50 de descuento = 700 + 10 % = 770 cobrado. Los
        # reportes por documento muestran 770; productos más vendidos
        # sigue en bruto (750), porque el descuento es del documento.
        self.sale.discount_type = "AMOUNT"
        self.sale.discount_value = Decimal("50")
        self.sale.tax_percentage = Decimal("10")
        self.sale.save(
            update_fields=["discount_type", "discount_value", "tax_percentage"],
        )
        freeze_sale_amounts(self.sale)

        params = {"date_from": "2026-07-01", "date_to": "2026-07-31"}

        by_date = self.client.get("/api/reports/sales-summary/", params)
        by_customer = self.client.get("/api/reports/top-customers/", params)
        by_product = self.client.get("/api/reports/top-selling-products/", params)

        self.assertEqual(
            by_date.data["summary"]["sales"]["total"],
            Decimal("770.00"),
        )
        self.assertEqual(
            Decimal(by_customer.data["results"][0]["total"]),
            Decimal("770.00"),
        )
        self.assertEqual(
            Decimal(by_product.data["results"][0]["total"]),
            Decimal("750.0000"),
        )

    def test_sales_reports_exclude_inactive_sales_like_the_cash_closing(self):
        # Punto 20: misma definición que la caja (confirmadas y activas).
        self.sale.is_active = False
        self.sale.save(update_fields=["is_active"])

        params = {"date_from": "2026-07-01", "date_to": "2026-07-31"}

        by_product = self.client.get("/api/reports/top-selling-products/", params)
        by_customer = self.client.get("/api/reports/top-customers/", params)

        self.assertEqual(by_product.data["count"], 0)
        self.assertEqual(by_customer.data["count"], 0)

    def test_top_selling_products_report(self):
        response = self.client.get(
            "/api/reports/top-selling-products/",
            {
                "date_from": "2026-07-01",
                "date_to": "2026-07-31",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data["results"]), 1)

        item = response.data["results"][0]

        self.assertEqual(
            item["product"]["standard_code"],
            "TOP-001",
        )
        self.assertEqual(item["quantity_sold"], 3)
        self.assertEqual(
            Decimal(item["total"]),
            Decimal("750.0000"),
        )

    def test_business_reports_reject_invalid_date(self):
        response = self.client.get(
            "/api/reports/sales-summary/",
            {
                "date_from": "fecha-mala",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_business_reports_reject_inverted_date_range(self):
        response = self.client.get(
            "/api/reports/sales-summary/",
            {
                "date_from": "2026-07-31",
                "date_to": "2026-07-01",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_purchases_by_supplier_report_converts_mixed_currencies_to_crc(self):
        second_purchase = Purchase.objects.create(
            supplier=self.supplier,
            invoice_number="REP-FAC-002",
            purchase_date=date(2026, 7, 15),
            currency="USD",
            exchange_rate=Decimal("520.0000"),
            status=PurchaseStatus.CONFIRMED,
            created_by=self.user,
            updated_by=self.user,
        )

        PurchaseItem.objects.create(
            purchase=second_purchase,
            supplier_product=self.supplier_product,
            quantity=2,
            unit_cost=Decimal("10.0000"),
            created_by=self.user,
            updated_by=self.user,
        )

        response = self.client.get(
            "/api/reports/purchases-by-supplier/",
            {
                "date_from": "2026-07-01",
                "date_to": "2026-07-31",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data["results"]), 1)

        item = response.data["results"][0]

        # Primera compra (setUp): 4 x 100.0000 CRC = 400.0000 CRC, sin conversión.
        # Segunda compra: 2 x 10.0000 USD = 20.0000 USD, convertidos con su
        # propio tipo de cambio: 20.0000 x 520.0000 = 10400.0000 CRC.
        # Total esperado: 400.0000 + 10400.0000 = 10800.0000 CRC.
        self.assertEqual(item["purchase_count"], 2)
        self.assertEqual(item["currency"], "CRC")
        self.assertEqual(
            Decimal(item["invoice_subtotal"]),
            Decimal("10800.0000"),
        )

    def test_business_reports_block_sales_only_user(self):
        # SALES ya no incluye view_reports en ROLE_PERMISSIONS.
        sales_user = User.objects.create_user(
            username="reports-sales-only",
            password="12345678",
        )

        sales_user.groups.add(
            Group.objects.get(name=ROLE_SALES),
        )

        self.client.force_authenticate(sales_user)

        purchases_response = self.client.get(
            "/api/reports/purchases-by-supplier/"
        )
        sales_response = self.client.get(
            "/api/reports/sales-summary/"
        )
        top_selling_response = self.client.get(
            "/api/reports/top-selling-products/"
        )

        self.assertEqual(purchases_response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(sales_response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(top_selling_response.status_code, status.HTTP_403_FORBIDDEN)

    def test_business_reports_block_inventory_only_user(self):
        # INVENTORY tampoco incluye view_reports en ROLE_PERMISSIONS.
        inventory_user = User.objects.create_user(
            username="reports-inventory-only",
            password="12345678",
        )

        inventory_user.groups.add(
            Group.objects.get(name=ROLE_INVENTORY),
        )

        self.client.force_authenticate(inventory_user)

        response = self.client.get(
            "/api/reports/purchases-by-supplier/"
        )

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_purchases_by_supplier_report_is_paginated(self):
        second_supplier = Supplier.objects.create(
            name="Proveedor Reportes Dos",
            created_by=self.user,
            updated_by=self.user,
        )

        second_supplier_product = SupplierProduct.objects.create(
            supplier=second_supplier,
            product=self.product,
            created_by=self.user,
            updated_by=self.user,
        )

        second_purchase = Purchase.objects.create(
            supplier=second_supplier,
            invoice_number="REP-FAC-010",
            purchase_date=date(2026, 7, 10),
            currency="CRC",
            status=PurchaseStatus.CONFIRMED,
            created_by=self.user,
            updated_by=self.user,
        )

        PurchaseItem.objects.create(
            purchase=second_purchase,
            supplier_product=second_supplier_product,
            quantity=1,
            unit_cost=Decimal("50.0000"),
            created_by=self.user,
            updated_by=self.user,
        )

        first_page = self.client.get(
            "/api/reports/purchases-by-supplier/",
            {
                "date_from": "2026-07-01",
                "date_to": "2026-07-31",
                "page_size": 1,
            },
        )

        self.assertEqual(first_page.status_code, status.HTTP_200_OK)
        self.assertEqual(first_page.data["count"], 2)
        self.assertEqual(len(first_page.data["results"]), 1)
        self.assertIsNotNone(first_page.data["next"])
        self.assertEqual(first_page.data["date_from"], date(2026, 7, 1))

        second_page = self.client.get(
            "/api/reports/purchases-by-supplier/",
            {
                "date_from": "2026-07-01",
                "date_to": "2026-07-31",
                "page_size": 1,
                "page": 2,
            },
        )

        self.assertEqual(second_page.status_code, status.HTTP_200_OK)
        self.assertEqual(len(second_page.data["results"]), 1)
        self.assertIsNone(second_page.data["next"])

    def test_top_selling_products_report_is_paginated(self):
        second_product = Product.objects.create(
            standard_code="TOP-002",
            name="Segundo producto vendido",
            storage_location=self.location,
            created_by=self.user,
            updated_by=self.user,
        )

        SaleItem.objects.create(
            sale=self.sale,
            product=second_product,
            quantity=1,
            unit_price=Decimal("100.0000"),
            created_by=self.user,
            updated_by=self.user,
        )

        first_page = self.client.get(
            "/api/reports/top-selling-products/",
            {
                "date_from": "2026-07-01",
                "date_to": "2026-07-31",
                "page_size": 1,
            },
        )

        self.assertEqual(first_page.status_code, status.HTTP_200_OK)
        self.assertEqual(first_page.data["count"], 2)
        self.assertEqual(len(first_page.data["results"]), 1)
        self.assertIsNotNone(first_page.data["next"])

        second_page = self.client.get(
            "/api/reports/top-selling-products/",
            {
                "date_from": "2026-07-01",
                "date_to": "2026-07-31",
                "page_size": 1,
                "page": 2,
            },
        )

        self.assertEqual(second_page.status_code, status.HTTP_200_OK)
        self.assertEqual(len(second_page.data["results"]), 1)
        self.assertIsNone(second_page.data["next"])

    def test_top_customers_report_orders_by_total_by_default(self):
        second_customer = Customer.objects.create(
            display_name="Cliente Frecuente",
            phone="8888-2222",
            created_by=self.user,
            updated_by=self.user,
        )

        first_frequent_sale = Sale.objects.create(
            customer=second_customer,
            sale_date=date(2026, 7, 5),
            currency="CRC",
            status=SaleStatus.CONFIRMED,
            created_by=self.user,
            updated_by=self.user,
        )

        SaleItem.objects.create(
            sale=first_frequent_sale,
            product=self.product,
            quantity=1,
            unit_price=Decimal("50.0000"),
            created_by=self.user,
            updated_by=self.user,
        )

        second_frequent_sale = Sale.objects.create(
            customer=second_customer,
            sale_date=date(2026, 7, 6),
            currency="CRC",
            status=SaleStatus.CONFIRMED,
            created_by=self.user,
            updated_by=self.user,
        )

        SaleItem.objects.create(
            sale=second_frequent_sale,
            product=self.product,
            quantity=1,
            unit_price=Decimal("50.0000"),
            created_by=self.user,
            updated_by=self.user,
        )

        by_total_response = self.client.get(
            "/api/reports/top-customers/",
            {
                "date_from": "2026-07-01",
                "date_to": "2026-07-31",
            },
        )

        self.assertEqual(by_total_response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(by_total_response.data["results"]), 2)

        # Por defecto ordena por monto total: Cliente Reportes (750.0000
        # en una sola venta) va antes que Cliente Frecuente (100.0000 en
        # dos ventas).
        first_by_total = by_total_response.data["results"][0]
        self.assertEqual(
            first_by_total["customer"]["display_name"],
            "CLIENTE REPORTES",
        )
        self.assertEqual(first_by_total["sale_count"], 1)
        self.assertEqual(
            Decimal(first_by_total["total"]),
            Decimal("750.0000"),
        )

        by_sale_count_response = self.client.get(
            "/api/reports/top-customers/",
            {
                "date_from": "2026-07-01",
                "date_to": "2026-07-31",
                "ordering": "sale_count",
            },
        )

        self.assertEqual(by_sale_count_response.status_code, status.HTTP_200_OK)

        # Ordenando por cantidad de ventas, Cliente Frecuente (2 ventas)
        # pasa a estar antes que Cliente Reportes (1 venta), aunque su
        # monto total sea menor.
        first_by_sale_count = by_sale_count_response.data["results"][0]
        self.assertEqual(
            first_by_sale_count["customer"]["display_name"],
            "CLIENTE FRECUENTE",
        )
        self.assertEqual(first_by_sale_count["sale_count"], 2)

    def test_top_customers_report_rejects_invalid_ordering(self):
        response = self.client.get(
            "/api/reports/top-customers/",
            {"ordering": "invalid-value"},
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_top_customers_report_excludes_sales_without_customer(self):
        anonymous_sale = Sale.objects.create(
            customer=None,
            sale_date=date(2026, 7, 7),
            currency="CRC",
            status=SaleStatus.CONFIRMED,
            created_by=self.user,
            updated_by=self.user,
        )

        SaleItem.objects.create(
            sale=anonymous_sale,
            product=self.product,
            quantity=1,
            unit_price=Decimal("999.0000"),
            created_by=self.user,
            updated_by=self.user,
        )

        response = self.client.get(
            "/api/reports/top-customers/",
            {
                "date_from": "2026-07-01",
                "date_to": "2026-07-31",
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data["results"]), 1)
        self.assertEqual(
            response.data["results"][0]["customer"]["display_name"],
            "CLIENTE REPORTES",
        )

    def test_top_customers_report_is_paginated(self):
        second_customer = Customer.objects.create(
            display_name="Cliente Reportes Dos",
            phone="8888-3333",
            created_by=self.user,
            updated_by=self.user,
        )

        second_customer_sale = Sale.objects.create(
            customer=second_customer,
            sale_date=date(2026, 7, 8),
            currency="CRC",
            status=SaleStatus.CONFIRMED,
            created_by=self.user,
            updated_by=self.user,
        )

        SaleItem.objects.create(
            sale=second_customer_sale,
            product=self.product,
            quantity=1,
            unit_price=Decimal("10.0000"),
            created_by=self.user,
            updated_by=self.user,
        )

        first_page = self.client.get(
            "/api/reports/top-customers/",
            {
                "date_from": "2026-07-01",
                "date_to": "2026-07-31",
                "page_size": 1,
            },
        )

        self.assertEqual(first_page.status_code, status.HTTP_200_OK)
        self.assertEqual(first_page.data["count"], 2)
        self.assertEqual(len(first_page.data["results"]), 1)
        self.assertIsNotNone(first_page.data["next"])

        second_page = self.client.get(
            "/api/reports/top-customers/",
            {
                "date_from": "2026-07-01",
                "date_to": "2026-07-31",
                "page_size": 1,
                "page": 2,
            },
        )

        self.assertEqual(second_page.status_code, status.HTTP_200_OK)
        self.assertEqual(len(second_page.data["results"]), 1)
        self.assertIsNone(second_page.data["next"])

    def test_product_supplier_prices_report_compares_suppliers_and_converts_currency(self):
        second_purchase = Purchase.objects.create(
            supplier=self.supplier,
            invoice_number="REP-FAC-003",
            purchase_date=date(2026, 7, 20),
            currency="CRC",
            status=PurchaseStatus.CONFIRMED,
            created_by=self.user,
            updated_by=self.user,
        )

        PurchaseItem.objects.create(
            purchase=second_purchase,
            supplier_product=self.supplier_product,
            quantity=2,
            unit_cost=Decimal("120.0000"),
            created_by=self.user,
            updated_by=self.user,
        )

        cheap_supplier = Supplier.objects.create(
            name="Proveedor Barato",
            created_by=self.user,
            updated_by=self.user,
        )

        cheap_supplier_product = SupplierProduct.objects.create(
            supplier=cheap_supplier,
            product=self.product,
            created_by=self.user,
            updated_by=self.user,
        )

        cheap_purchase = Purchase.objects.create(
            supplier=cheap_supplier,
            invoice_number="REP-FAC-004",
            purchase_date=date(2026, 7, 15),
            currency="USD",
            exchange_rate=Decimal("600.0000"),
            status=PurchaseStatus.CONFIRMED,
            created_by=self.user,
            updated_by=self.user,
        )

        PurchaseItem.objects.create(
            purchase=cheap_purchase,
            supplier_product=cheap_supplier_product,
            quantity=5,
            unit_cost=Decimal("0.1500"),
            created_by=self.user,
            updated_by=self.user,
        )

        response = self.client.get(
            "/api/reports/product-supplier-prices/",
            {
                "product": self.product.id,
            },
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response.data["product"]["standard_code"],
            "TOP-001",
        )
        self.assertEqual(len(response.data["results"]), 2)

        # Ordenado por precio unitario más reciente ascendente: el
        # proveedor más barato primero, aunque su compra en dólares
        # requiera conversión.
        cheapest = response.data["results"][0]
        self.assertEqual(cheapest["supplier"]["name"], "PROVEEDOR BARATO")
        self.assertEqual(cheapest["purchase_count"], 1)
        self.assertEqual(cheapest["currency"], "CRC")
        self.assertEqual(
            Decimal(cheapest["last_unit_cost"]),
            Decimal("90.0000"),
        )
        self.assertEqual(
            Decimal(cheapest["average_unit_cost"]),
            Decimal("90.0000"),
        )
        self.assertEqual(len(cheapest["purchases"]), 1)
        self.assertEqual(
            cheapest["purchases"][0]["invoice_number"],
            "REP-FAC-004",
        )
        self.assertEqual(
            Decimal(cheapest["purchases"][0]["unit_cost"]),
            Decimal("90.0000"),
        )

        # Proveedor original: dos compras en colones, sin conversión.
        # Último precio: 120.0000 (2026-07-20). Promedio: (100+120)/2=110.0000.
        other = response.data["results"][1]
        self.assertEqual(other["supplier"]["name"], "PROVEEDOR REPORTES")
        self.assertEqual(other["purchase_count"], 2)
        self.assertEqual(
            other["last_purchase_date"],
            date(2026, 7, 20),
        )
        self.assertEqual(
            Decimal(other["last_unit_cost"]),
            Decimal("120.0000"),
        )
        self.assertEqual(
            Decimal(other["average_unit_cost"]),
            Decimal("110.0000"),
        )

        # Historial de compras a este proveedor en orden cronológico,
        # para poder ver si el precio cambió con el tiempo.
        self.assertEqual(len(other["purchases"]), 2)

        first_purchase_entry = other["purchases"][0]
        self.assertEqual(first_purchase_entry["invoice_number"], "REP-FAC-001")
        self.assertEqual(first_purchase_entry["purchase_date"], date(2026, 7, 1))
        self.assertEqual(
            Decimal(first_purchase_entry["unit_cost"]),
            Decimal("100.0000"),
        )

        second_purchase_entry = other["purchases"][1]
        self.assertEqual(second_purchase_entry["invoice_number"], "REP-FAC-003")
        self.assertEqual(second_purchase_entry["purchase_date"], date(2026, 7, 20))
        self.assertEqual(
            Decimal(second_purchase_entry["unit_cost"]),
            Decimal("120.0000"),
        )

    def test_product_supplier_prices_report_requires_product(self):
        response = self.client.get(
            "/api/reports/product-supplier-prices/"
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_product_supplier_prices_report_rejects_unknown_product(self):
        response = self.client.get(
            "/api/reports/product-supplier-prices/",
            {
                "product": 999999,
            },
        )

        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)

    def test_product_supplier_prices_report_is_paginated(self):
        second_supplier = Supplier.objects.create(
            name="Proveedor Reportes Dos",
            created_by=self.user,
            updated_by=self.user,
        )

        second_supplier_product = SupplierProduct.objects.create(
            supplier=second_supplier,
            product=self.product,
            created_by=self.user,
            updated_by=self.user,
        )

        second_purchase = Purchase.objects.create(
            supplier=second_supplier,
            invoice_number="REP-FAC-005",
            purchase_date=date(2026, 7, 12),
            currency="CRC",
            status=PurchaseStatus.CONFIRMED,
            created_by=self.user,
            updated_by=self.user,
        )

        PurchaseItem.objects.create(
            purchase=second_purchase,
            supplier_product=second_supplier_product,
            quantity=1,
            unit_cost=Decimal("50.0000"),
            created_by=self.user,
            updated_by=self.user,
        )

        first_page = self.client.get(
            "/api/reports/product-supplier-prices/",
            {
                "product": self.product.id,
                "page_size": 1,
            },
        )

        self.assertEqual(first_page.status_code, status.HTTP_200_OK)
        self.assertEqual(first_page.data["count"], 2)
        self.assertEqual(len(first_page.data["results"]), 1)
        self.assertIsNotNone(first_page.data["next"])

        second_page = self.client.get(
            "/api/reports/product-supplier-prices/",
            {
                "product": self.product.id,
                "page_size": 1,
                "page": 2,
            },
        )

        self.assertEqual(second_page.status_code, status.HTTP_200_OK)
        self.assertEqual(len(second_page.data["results"]), 1)
        self.assertIsNone(second_page.data["next"])


class ComputeDocumentAmountsTest(TestCase):
    """
    Cálculo compartido de venta, servicio y proforma (puntos 15, 19 y
    26): descuento sobre el subtotal, impuesto sobre la base con
    descuento, redondeo a 2 decimales ROUND_HALF_UP.
    """

    def test_without_discount_or_tax_total_is_the_subtotal(self):
        amounts = compute_document_amounts(Decimal("750.0000"))

        self.assertEqual(amounts.subtotal, Decimal("750.00"))
        self.assertEqual(amounts.discount_amount, Decimal("0.00"))
        self.assertEqual(amounts.tax_amount, Decimal("0.00"))
        self.assertEqual(amounts.total, Decimal("750.00"))

    def test_amount_discount_then_tax_on_discounted_base(self):
        amounts = compute_document_amounts(
            Decimal("1000"),
            DiscountType.AMOUNT,
            Decimal("100"),
            Decimal("13"),
        )

        self.assertEqual(amounts.discount_amount, Decimal("100.00"))
        self.assertEqual(amounts.taxable_base, Decimal("900.00"))
        self.assertEqual(amounts.tax_amount, Decimal("117.00"))
        self.assertEqual(amounts.total, Decimal("1017.00"))

    def test_percent_discount_rounds_half_up(self):
        # 10 % de 33.35 = 3.335 → 3.34; base 30.01; 13 % = 3.9013 → 3.90.
        amounts = compute_document_amounts(
            Decimal("33.35"),
            DiscountType.PERCENT,
            Decimal("10"),
            Decimal("13"),
        )

        self.assertEqual(amounts.discount_amount, Decimal("3.34"))
        self.assertEqual(amounts.taxable_base, Decimal("30.01"))
        self.assertEqual(amounts.tax_amount, Decimal("3.90"))
        self.assertEqual(amounts.total, Decimal("33.91"))

    def test_full_discount_is_allowed(self):
        amounts = compute_document_amounts(
            Decimal("500"),
            DiscountType.PERCENT,
            Decimal("100"),
            Decimal("13"),
        )

        self.assertEqual(amounts.total, Decimal("0.00"))

    def test_invalid_values_raise(self):
        cases = {
            "descuento mayor al subtotal": (
                Decimal("100"),
                DiscountType.AMOUNT,
                Decimal("100.01"),
                Decimal("0"),
            ),
            "porcentaje mayor a 100": (
                Decimal("100"),
                DiscountType.PERCENT,
                Decimal("101"),
                Decimal("0"),
            ),
            "descuento negativo": (
                Decimal("100"),
                DiscountType.AMOUNT,
                Decimal("-1"),
                Decimal("0"),
            ),
            "impuesto mayor a 100": (
                Decimal("100"),
                DiscountType.AMOUNT,
                Decimal("0"),
                Decimal("101"),
            ),
            "tipo inválido": (
                Decimal("100"),
                "OTRO",
                Decimal("0"),
                Decimal("0"),
            ),
        }

        for case, arguments in cases.items():
            with self.subTest(case=case):
                with self.assertRaises(DocumentAmountsError):
                    compute_document_amounts(*arguments)


from datetime import datetime, time  # noqa: E402

from apps.customers.services import deliver_service  # noqa: E402


class SalesSummaryReportApiTest(APITestCase):
    """
    Punto 20: resumen de ventas y servicios por período, con el mismo
    criterio que la caja (confirmadas/entregados y activos, montos
    cobrados) y comparación contra otro rango.
    """

    def setUp(self):
        call_command("setup_roles")

        self.user = User.objects.create_user(
            username="summary-admin",
            password="12345678",
        )
        self.user.groups.add(Group.objects.get(name=ROLE_ADMIN))
        self.client.force_authenticate(self.user)

        location = StorageLocation.objects.create(
            code="S101",
            created_by=self.user,
            updated_by=self.user,
        )
        self.product = Product.objects.create(
            standard_code="SUM-001",
            name="Producto resumen",
            storage_location=location,
            created_by=self.user,
            updated_by=self.user,
        )
        customer = Customer.objects.create(
            display_name="Cliente resumen",
            created_by=self.user,
            updated_by=self.user,
        )
        self.injector = Injector.objects.create(
            customer=customer,
            injector_number="SUM-INY",
            created_by=self.user,
            updated_by=self.user,
        )

    def create_sale(
        self,
        sale_date,
        amount="1000.0000",
        payment_method="CASH",
        tax_percentage="0",
        **extra,
    ):
        sale = Sale.objects.create(
            sale_date=sale_date,
            currency="CRC",
            status=extra.pop("status", SaleStatus.CONFIRMED),
            payment_method=payment_method,
            tax_percentage=Decimal(tax_percentage),
            created_by=self.user,
            updated_by=self.user,
            **extra,
        )
        SaleItem.objects.create(
            sale=sale,
            product=self.product,
            quantity=1,
            unit_price=Decimal(amount),
            created_by=self.user,
            updated_by=self.user,
        )

        if sale.status == SaleStatus.CONFIRMED:
            freeze_sale_amounts(sale)

        return sale

    def create_service(
        self,
        delivered_on,
        price="20000.0000",
        payment_method="CARD",
        tax_percentage="0",
        is_active=True,
    ):
        service_record = InjectorServiceRecord.objects.create(
            injector=self.injector,
            received_at=timezone.now(),
            status=InjectorServiceStatus.READY,
            payment_method=payment_method,
            price=Decimal(price),
            tax_percentage=Decimal(tax_percentage),
            is_active=is_active,
            created_by=self.user,
            updated_by=self.user,
        )

        return deliver_service(
            service_record=service_record,
            delivered_at=timezone.make_aware(datetime.combine(delivered_on, time(10))),
            user=self.user,
        )

    def summary(self, **params):
        response = self.client.get(
            "/api/reports/sales-summary/",
            {"date_from": "2026-07-01", "date_to": "2026-07-31", **params},
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)

        return response.data

    def test_sales_services_and_business_total(self):
        # Venta 1 000 + 13 % = 1 130 en efectivo; servicio 20 000 + 13 %
        # = 22 600 con tarjeta.
        self.create_sale(date(2026, 7, 2), tax_percentage="13")
        self.create_service(date(2026, 7, 3), tax_percentage="13")

        summary = self.summary()["summary"]

        self.assertEqual(summary["sales"]["total"], Decimal("1130.00"))
        self.assertEqual(summary["services"]["total"], Decimal("22600.00"))
        self.assertEqual(summary["total"]["count"], 2)
        self.assertEqual(summary["total"]["total"], Decimal("23730.00"))
        self.assertEqual(summary["total"]["tax"], Decimal("2730.00"))
        self.assertEqual(summary["total"]["business_income"], Decimal("21000.00"))
        self.assertEqual(summary["total"]["by_method"]["CASH"], Decimal("1130.00"))
        self.assertEqual(summary["total"]["by_method"]["CARD"], Decimal("22600.00"))

    def test_weeks_run_saturday_to_friday_like_the_cash_closing(self):
        self.create_sale(date(2026, 7, 3))  # viernes
        self.create_sale(date(2026, 7, 4))  # sábado

        periods = self.summary(group_by="week")["periods"]

        self.assertEqual(
            [(p["period_start"], p["period_end"]) for p in periods],
            [
                (date(2026, 6, 27), date(2026, 7, 3)),
                (date(2026, 7, 4), date(2026, 7, 10)),
            ],
        )

    def test_months_group_by_calendar_month(self):
        self.create_sale(date(2026, 7, 31))
        self.create_sale(date(2026, 8, 1))

        periods = self.summary(date_to="2026-08-31", group_by="month")["periods"]

        self.assertEqual(
            [(p["period_start"], p["period_end"]) for p in periods],
            [
                (date(2026, 7, 1), date(2026, 7, 31)),
                (date(2026, 8, 1), date(2026, 8, 31)),
            ],
        )

    def test_excludes_what_the_cash_closing_excludes(self):
        self.create_sale(date(2026, 7, 2))
        self.create_sale(date(2026, 7, 2), status=SaleStatus.DRAFT)
        self.create_sale(date(2026, 7, 2), status=SaleStatus.CANCELLED)
        self.create_sale(date(2026, 7, 2), is_active=False)
        self.create_service(date(2026, 7, 2), is_active=False)

        summary = self.summary()["summary"]

        self.assertEqual(summary["sales"]["count"], 1)
        self.assertEqual(summary["services"]["count"], 0)
        self.assertEqual(summary["total"]["total"], Decimal("1000.00"))

    def test_comparison_returns_previous_summary_and_variation(self):
        self.create_sale(date(2026, 6, 10), amount="500.0000")
        self.create_sale(date(2026, 7, 10), amount="1000.0000")

        data = self.summary(compare_from="2026-06-01", compare_to="2026-06-30")
        comparison = data["comparison"]

        self.assertEqual(comparison["summary"]["total"]["total"], Decimal("500.00"))
        self.assertEqual(
            comparison["variation"]["total"]["total"],
            {"amount": Decimal("500.00"), "percentage": Decimal("100.0")},
        )
        # Sin servicios en ninguno de los dos: crecer desde 0 no es un %.
        self.assertIsNone(
            comparison["variation"]["services"]["total"]["percentage"],
        )

    def test_rejects_invalid_parameters(self):
        cases = {
            "sin fechas": {},
            "group_by inválido": {
                "date_from": "2026-07-01",
                "date_to": "2026-07-31",
                "group_by": "year",
            },
            "comparación incompleta": {
                "date_from": "2026-07-01",
                "date_to": "2026-07-31",
                "compare_from": "2026-06-01",
            },
        }

        for case, params in cases.items():
            with self.subTest(case=case):
                response = self.client.get("/api/reports/sales-summary/", params)

                self.assertEqual(
                    response.status_code,
                    status.HTTP_400_BAD_REQUEST,
                )


from unittest import mock  # noqa: E402

from reportlab import rl_config  # noqa: E402


def pdf_text(response):
    """
    Contenido del PDF como texto. Las pruebas lo generan sin comprimir
    (rl_config.pageCompression = 0), así el texto ASCII queda legible.
    """

    return b"".join(response.streaming_content).decode("latin-1")


@mock.patch.object(rl_config, "pageCompression", 0)
class ReportPdfApiTest(APITestCase):
    """Puntos 2 + 27: PDF de cada reporte, con sus filtros y su permiso."""

    REPORT_PDF_URLS = {
        "bajo mínimo": ("/api/reports/low-stock-products/pdf/", {}),
        "stock por ubicación": ("/api/reports/stock-by-location/pdf/", {}),
        "compras por proveedor": ("/api/reports/purchases-by-supplier/pdf/", {}),
        "resumen de ventas": (
            "/api/reports/sales-summary/pdf/",
            {"date_from": "2026-07-01", "date_to": "2026-07-31"},
        ),
        "más vendidos": ("/api/reports/top-selling-products/pdf/", {}),
        "clientes": ("/api/reports/top-customers/pdf/", {}),
    }

    setUp_summary = SalesSummaryReportApiTest.setUp
    create_sale = SalesSummaryReportApiTest.create_sale

    def setUp(self):
        self.setUp_summary()

        self.read_only_user = User.objects.create_user(
            username="pdf-readonly",
            password="12345678",
        )
        self.read_only_user.groups.add(Group.objects.get(name=ROLE_READ_ONLY))

        self.location = StorageLocation.objects.create(
            code="PDF-A",
            description="Estante PDF",
            created_by=self.user,
            updated_by=self.user,
        )
        self.low_product = Product.objects.create(
            standard_code="PDF-LOW",
            name="Producto bajo",
            storage_location=self.location,
            minimum_stock=5,
            created_by=self.user,
            updated_by=self.user,
        )
        self.stocked_product = Product.objects.create(
            standard_code="PDF-OK",
            name="Producto con stock",
            storage_location=self.location,
            minimum_stock=1,
            created_by=self.user,
            updated_by=self.user,
        )
        initial_inventory(product=self.stocked_product, quantity=10, user=self.user)

    def get_pdf(self, url, params=None):
        response = self.client.get(url, params or {})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response["Content-Type"], "application/pdf")

        text = pdf_text(response)
        self.assertTrue(text.startswith("%PDF"))

        return text

    def test_every_report_has_a_pdf_with_the_signature_box(self):
        urls = dict(self.REPORT_PDF_URLS)
        urls["precios por proveedor"] = (
            "/api/reports/product-supplier-prices/pdf/",
            {"product": self.stocked_product.id},
        )

        for case, (url, params) in urls.items():
            with self.subTest(case=case):
                text = self.get_pdf(url, params)

                self.assertIn("(Revisado por:) Tj", text)
                self.assertIn("(Fecha:) Tj", text)
                self.assertIn("(Observaciones:) Tj", text)
                self.assertIn("gina 1 de 1", text)

    def test_report_pdfs_require_view_reports(self):
        self.client.force_authenticate(self.read_only_user)

        for case, (url, params) in self.REPORT_PDF_URLS.items():
            with self.subTest(case=case):
                response = self.client.get(url, params)

                self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_low_stock_pdf_has_a_column_to_write_the_quantity_to_order(self):
        text = self.get_pdf("/api/reports/low-stock-products/pdf/")

        self.assertIn("(PDF-LOW) Tj", text)
        self.assertNotIn("(PDF-OK) Tj", text)
        self.assertIn("(CANTIDAD A PEDIR) Tj", text)

    def test_stock_by_location_pdf_respects_filters(self):
        text = self.get_pdf(
            "/api/reports/stock-by-location/pdf/",
            {"location": self.location.id, "q": "PDF-OK"},
        )

        self.assertIn("(PDF-OK) Tj", text)
        self.assertNotIn("(PDF-LOW) Tj", text)
        self.assertIn("(CONTEO F", text)
        self.assertIn("n: PDF-A.", text)

    def test_long_table_continues_on_the_next_page_with_its_header(self):
        for number in range(45):
            Product.objects.create(
                standard_code=f"PDF-{number:03d}",
                name=f"Producto {number}",
                storage_location=self.location,
                created_by=self.user,
                updated_by=self.user,
            )

        text = self.get_pdf(
            "/api/reports/stock-by-location/pdf/",
            {"location": self.location.id},
        )

        self.assertIn("gina 2 de", text)
        self.assertGreaterEqual(text.count("(CONTEO F"), 2)
        self.assertIn("(PDF-044) Tj", text)

    def test_sales_reports_pdf_list_the_sales_of_the_range(self):
        self.create_sale(date(2026, 7, 10), amount="2500.0000")
        self.create_sale(date(2026, 8, 10), amount="9900.0000")

        july = {"date_from": "2026-07-01", "date_to": "2026-07-31"}

        text = self.get_pdf("/api/reports/top-selling-products/pdf/", july)
        self.assertIn("(SUM-001) Tj", text)
        self.assertIn("(2 500) Tj", text)
        self.assertNotIn("(9 900) Tj", text)

        text = self.get_pdf("/api/reports/sales-summary/pdf/", july)
        self.assertIn("(10/07/2026) Tj", text)
        self.assertNotIn("Comparaci", text)

        text = self.get_pdf(
            "/api/reports/sales-summary/pdf/",
            {
                **july,
                "compare_from": "2026-08-01",
                "compare_to": "2026-08-31",
            },
        )
        self.assertIn("Comparaci", text)
        self.assertIn("(-7 400) Tj", text)

    def test_report_pdfs_validate_their_filters(self):
        cases = {
            "fecha inválida": (
                "/api/reports/purchases-by-supplier/pdf/",
                {"date_from": "01/07/2026"},
            ),
            "rango invertido": (
                "/api/reports/top-selling-products/pdf/",
                {"date_from": "2026-08-01", "date_to": "2026-07-01"},
            ),
            "orden inválido": (
                "/api/reports/top-customers/pdf/",
                {"ordering": "name"},
            ),
            "resumen sin fechas": ("/api/reports/sales-summary/pdf/", {}),
            "precios sin producto": (
                "/api/reports/product-supplier-prices/pdf/",
                {},
            ),
            "ubicación inválida": (
                "/api/reports/stock-by-location/pdf/",
                {"location": "abc"},
            ),
        }

        for case, (url, params) in cases.items():
            with self.subTest(case=case):
                response = self.client.get(url, params)

                self.assertEqual(
                    response.status_code,
                    status.HTTP_400_BAD_REQUEST,
                )


class ImpossibleDateFilterTest(APITestCase):
    """
    Una fecha bien escrita que no existe (2026-02-30, 2026-13-01) es un
    error del filtro: 400, no 500. parse_date de Django lanza ValueError
    en ese caso; los parsers la tratan como formato inválido.
    """

    def setUp(self):
        call_command("setup_roles")

        user = User.objects.create_user(
            username="impossible-dates-admin",
            password="12345678",
        )
        user.groups.add(Group.objects.get(name=ROLE_ADMIN))
        self.client.force_authenticate(user)

    def test_impossible_dates_return_400(self):
        cases = {
            "movimientos": (
                "/api/inventory/stock-movements/",
                {"date_from": "2026-02-30"},
            ),
            "PDF de movimientos": (
                "/api/inventory/stock-movements/pdf/",
                {"date_to": "2026-13-01"},
            ),
            "cierres de caja": (
                "/api/cash/closings/",
                {"date_from": "2026-02-30"},
            ),
            "compras por proveedor": (
                "/api/reports/purchases-by-supplier/",
                {"date_from": "2026-13-01"},
            ),
            "más vendidos": (
                "/api/reports/top-selling-products/",
                {"date_to": "2026-02-30"},
            ),
            "clientes (PDF)": (
                "/api/reports/top-customers/pdf/",
                {"date_from": "2026-02-30"},
            ),
            "resumen de ventas": (
                "/api/reports/sales-summary/",
                {"date_from": "2026-02-30", "date_to": "2026-03-31"},
            ),
            "comparación del resumen": (
                "/api/reports/sales-summary/pdf/",
                {
                    "date_from": "2026-03-01",
                    "date_to": "2026-03-31",
                    "compare_from": "2026-02-01",
                    "compare_to": "2026-02-30",
                },
            ),
        }

        for case, (url, params) in cases.items():
            with self.subTest(case=case):
                response = self.client.get(url, params)

                self.assertEqual(
                    response.status_code,
                    status.HTTP_400_BAD_REQUEST,
                )

    def test_valid_dates_still_work(self):
        response = self.client.get(
            "/api/reports/top-selling-products/",
            {"date_from": "2024-02-29", "date_to": "2024-03-31"},
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)


from io import BytesIO  # noqa: E402

from openpyxl import load_workbook  # noqa: E402

from apps.inventory.models import (  # noqa: E402
    MovementDirection,
    StockMovement,
    StockMovementType,
)


def xlsx_rows(response):
    """Filas de la tabla del Excel (desde el encabezado) como tuplas."""

    workbook = load_workbook(BytesIO(b"".join(response.streaming_content)))
    rows = list(workbook.active.iter_rows(values_only=True))
    header_index = next(
        index for index, row in enumerate(rows) if row[0] in ("#", "Fecha")
    )

    return rows[header_index], rows[header_index + 1:]


class MostMovedProductsReportApiTest(APITestCase):
    """
    Punto 28: productos que más se mueven. Neto de reversiones y sin
    inventario inicial; ajustes y conteos cuentan.
    """

    URL = "/api/reports/most-moved-products/"

    def setUp(self):
        call_command("setup_roles")

        self.user = User.objects.create_user(
            username="most-moved-admin",
            password="12345678",
        )
        self.user.groups.add(Group.objects.get(name=ROLE_ADMIN))
        self.client.force_authenticate(self.user)

        location = StorageLocation.objects.create(
            code="MOV-A",
            created_by=self.user,
            updated_by=self.user,
        )

        def product(code):
            return Product.objects.create(
                standard_code=code,
                name=f"Producto {code}",
                storage_location=location,
                created_by=self.user,
                updated_by=self.user,
            )

        self.busy = product("MOV-BUSY")
        self.reversed = product("MOV-REV")
        self.initial_only = product("MOV-INIT")

    def move(self, product, movement_type, direction, quantity, reverses=None):
        return StockMovement.create_from_service(
            product=product,
            movement_type=movement_type,
            direction=direction,
            quantity=quantity,
            notes="Prueba",
            reverses_movement=reverses,
            created_by=self.user,
            updated_by=self.user,
        )

    def create_movements(self):
        IN, OUT = MovementDirection.IN, MovementDirection.OUT
        # Ajustes: una entrada o salida "de verdad" exige su línea de
        # compra o de venta, y al ranking solo le importa la dirección.
        ADJUSTMENT = StockMovementType.ADJUSTMENT

        # 10 de entrada y 5 de salida; el inventario inicial no cuenta.
        self.move(self.busy, StockMovementType.INITIAL, IN, 100)
        self.move(self.busy, ADJUSTMENT, IN, 10)
        self.move(self.busy, ADJUSTMENT, OUT, 4)
        self.move(self.busy, ADJUSTMENT, OUT, 1)

        # Una salida de 20 que se revirtió no cuenta, ni su reversión;
        # queda solo la salida de 2.
        cancelled = self.move(self.reversed, ADJUSTMENT, OUT, 20)
        self.move(
            self.reversed,
            StockMovementType.REVERSAL,
            IN,
            20,
            reverses=cancelled,
        )
        self.move(self.reversed, ADJUSTMENT, OUT, 2)

        self.move(self.initial_only, StockMovementType.INITIAL, IN, 50)

    def test_ranking_is_net_of_reversals_and_skips_initial_inventory(self):
        self.create_movements()

        response = self.client.get(self.URL)

        self.assertEqual(response.status_code, status.HTTP_200_OK)

        results = response.data["results"]

        self.assertEqual(
            [item["product"]["standard_code"] for item in results],
            ["MOV-BUSY", "MOV-REV"],
        )
        self.assertEqual(
            {
                key: results[0][key]
                for key in (
                    "quantity_in",
                    "quantity_out",
                    "total_moved",
                    "movement_count",
                )
            },
            {
                "quantity_in": 10,
                "quantity_out": 5,
                "total_moved": 15,
                "movement_count": 3,
            },
        )
        self.assertEqual(results[0]["location_code"], "MOV-A")
        self.assertEqual(results[1]["total_moved"], 2)
        self.assertEqual(results[1]["movement_count"], 1)

    def test_ranking_respects_the_date_range(self):
        self.create_movements()

        today = timezone.localdate().isoformat()

        response = self.client.get(self.URL, {"date_from": today, "date_to": today})
        self.assertEqual(len(response.data["results"]), 2)

        response = self.client.get(self.URL, {"date_to": "2000-01-01"})
        self.assertEqual(response.data["results"], [])

        response = self.client.get(self.URL, {"date_from": "2026/01/01"})
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_excel_has_the_same_rows(self):
        self.create_movements()

        response = self.client.get(f"{self.URL}excel/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        header, rows = xlsx_rows(response)

        self.assertEqual(
            header,
            (
                "#",
                "Código",
                "Producto",
                "Ubicación",
                "Entradas",
                "Salidas",
                "Total movido",
                "Movimientos",
            ),
        )
        self.assertEqual(
            rows,
            [
                (1, "MOV-BUSY", "Producto MOV-BUSY", "MOV-A", 10, 5, 15, 3),
                (2, "MOV-REV", "Producto MOV-REV", "MOV-A", 0, 2, 2, 1),
            ],
        )

    def test_pdf_is_generated(self):
        self.create_movements()

        response = self.client.get(f"{self.URL}pdf/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response["Content-Type"], "application/pdf")

    def test_requires_view_reports(self):
        read_only_user = User.objects.create_user(
            username="most-moved-readonly",
            password="12345678",
        )
        read_only_user.groups.add(Group.objects.get(name=ROLE_READ_ONLY))
        self.client.force_authenticate(read_only_user)

        for suffix in ("", "pdf/", "excel/"):
            with self.subTest(url=f"{self.URL}{suffix}"):
                response = self.client.get(f"{self.URL}{suffix}")

                self.assertEqual(
                    response.status_code,
                    status.HTTP_403_FORBIDDEN,
                )
